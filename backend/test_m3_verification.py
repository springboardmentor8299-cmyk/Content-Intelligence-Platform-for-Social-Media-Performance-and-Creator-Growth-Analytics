import io
import openpyxl
from fastapi.testclient import TestClient
from app.main import app
from app.integrations import (
    YouTubeIntegrationService,
    InstagramIntegrationService,
    FacebookIntegrationService,
    XIntegrationService,
    LinkedInIntegrationService
)

client = TestClient(app)

def run_tests():
    print("==================================================")
    print("CREATORIQ — MILESTONE 3 COMPREHENSIVE VERIFICATION")
    print("==================================================")
    
    # 1. Root & Health
    res = client.get("/")
    assert res.status_code == 200, f"Root failed: {res.status_code}"
    print("[PASS] 1. FastAPI App Startup & Health Endpoint")

    # 2. Authentication & JWT
    login_res = client.post("/api/v1/auth/login", json={
        "email": "creator@creatoriq.io",
        "password": "Creator@123"
    })
    assert login_res.status_code == 200, f"Login failed: {login_res.status_code}"
    token_data = login_res.json()
    token = token_data["access_token"]
    auth_headers = {"Authorization": f"Bearer {token}"}
    print("[PASS] 2. User Authentication & JWT Bearer Token")

    # 3. RBAC Persona Switching
    for role in ["agency", "marketing_team", "admin", "creator"]:
        switch_res = client.post(f"/api/v1/auth/switch-role/{role}", headers=auth_headers)
        assert switch_res.status_code == 200, f"Role switch to {role} failed"
    print("[PASS] 3. RBAC Role Switching (creator, agency, marketing_team, admin)")

    # 4. Revenue / Monetization CRUD & Analysis
    rev_list = client.get("/api/v1/revenue", headers=auth_headers)
    assert rev_list.status_code == 200
    records = rev_list.json()
    assert len(records) > 0, "No revenue records found"

    summary_res = client.get("/api/v1/revenue/summary", headers=auth_headers)
    assert summary_res.status_code == 200
    summary = summary_res.json()
    assert summary["total_revenue"] > 0
    assert summary["total_paid"] > 0
    assert summary["total_pending"] > 0
    assert "source_breakdown" in summary
    assert "disclaimer" in summary
    print(f"[PASS] 4a. Revenue Summary (${summary['total_revenue']:,.2f} total across {summary['deals_count']} deals)")

    # Create new deal
    new_deal_res = client.post("/api/v1/revenue", headers=auth_headers, json={
        "title": "M3 Verification Test Deal",
        "brand_name": "Test Brand Sponsor",
        "amount": 2500.0,
        "source": "sponsorship",
        "status": "pending",
        "notes": "Automated verification test item"
    })
    assert new_deal_res.status_code == 201
    created_deal = new_deal_res.json()
    created_id = created_deal["id"]
    print(f"[PASS] 4b. Revenue Deal Creation (ID: {created_id})")

    # Update deal
    update_res = client.put(f"/api/v1/revenue/{created_id}", headers=auth_headers, json={
        "status": "paid",
        "amount": 2750.0
    })
    assert update_res.status_code == 200
    assert update_res.json()["status"] == "paid"
    assert update_res.json()["amount"] == 2750.0
    print("[PASS] 4c. Revenue Deal Update")

    # Delete deal
    del_res = client.delete(f"/api/v1/revenue/{created_id}", headers=auth_headers)
    assert del_res.status_code == 200
    print("[PASS] 4d. Revenue Deal Deletion")

    # 5. Notifications
    notifs_res = client.get("/api/v1/notifications", headers=auth_headers)
    assert notifs_res.status_code == 200
    notifs = notifs_res.json()
    assert len(notifs) > 0

    unread_res = client.get("/api/v1/notifications/unread-count", headers=auth_headers)
    assert unread_res.status_code == 200
    unread_count = unread_res.json()["unread_count"]

    first_notif_id = notifs[0]["id"]
    mark_res = client.patch(f"/api/v1/notifications/{first_notif_id}/read", headers=auth_headers)
    assert mark_res.status_code == 200
    assert mark_res.json()["is_read"] is True

    mark_all_res = client.patch("/api/v1/notifications/read-all", headers=auth_headers)
    assert mark_all_res.status_code == 200
    print(f"[PASS] 5. Notifications Listing, Unread Count ({unread_count}), Read & Read-All")

    # 6. Reports & Export (All 6 Reports x 4 Periods)
    report_types = [
        "analytics_summary",
        "content_performance",
        "audience_analytics",
        "growth_trends",
        "revenue",
        "platform_comparison"
    ]
    periods = ["7d", "30d", "90d", "all_time"]

    for rt in report_types:
        for p in periods:
            prev_res = client.get(f"/api/v1/reports/preview?report_type={rt}&period={p}", headers=auth_headers)
            assert prev_res.status_code == 200, f"Preview failed for {rt} ({p})"
            data = prev_res.json()
            assert "kpis" in data
            assert "data_provenance" in data
            if rt == "platform_comparison":
                platforms = [i["platform"] for i in data.get("items", [])]
                assert "YouTube" in platforms
                assert "Instagram" in platforms
                assert "Facebook" in platforms
                assert "X" in platforms
                assert "LinkedIn" in platforms
                assert "TikTok" not in platforms
    print("[PASS] 6a. All 6 Reports Previewed across all 4 periods (7d, 30d, 90d, all_time)")

    # Test CSV Export
    csv_res = client.get("/api/v1/reports/export?report_type=platform_comparison&period=30d&format=csv", headers=auth_headers)
    assert csv_res.status_code == 200
    assert "text/csv" in csv_res.headers.get("content-type", "")
    assert len(csv_res.content) > 100
    assert b"TikTok" not in csv_res.content
    assert b"YouTube" in csv_res.content
    assert b"X" in csv_res.content
    print("[PASS] 6b. CSV Report Export verified")

    # Test XLSX Export
    xlsx_res = client.get("/api/v1/reports/export?report_type=revenue&period=30d&format=xlsx", headers=auth_headers)
    assert xlsx_res.status_code == 200
    assert "spreadsheetml" in xlsx_res.headers.get("content-type", "")
    # Verify openpyxl can load the generated binary bytes
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_res.content))
    assert "Executive Report" in wb.sheetnames
    ws = wb["Executive Report"]
    assert ws.cell(row=2, column=1).value is not None
    print("[PASS] 6c. Real Excel (XLSX) Export verified with openpyxl binary inspection")

    # 7. KPI Monitoring Endpoint
    kpi_res = client.get("/api/v1/analytics/kpi-summary", headers=auth_headers)
    assert kpi_res.status_code == 200
    kpi_data = kpi_res.json()
    assert kpi_data["content_performance"]["monitored_items"] == 26
    assert kpi_data["audience_status"]["total_subscribers"] == 1420000
    assert kpi_data["social_integration_status"]["total_supported"] == 5
    assert "youtube" in kpi_data["social_integration_status"]["platforms"]
    assert "x" in kpi_data["social_integration_status"]["platforms"]
    assert "tiktok" not in kpi_data["social_integration_status"]["platforms"]
    print("[PASS] 7. KPI Monitoring across all 8 project dimensions (5 platforms)")

    # 8. Social Accounts & Integrations
    accounts_res = client.get("/api/v1/social/accounts", headers=auth_headers)
    assert accounts_res.status_code == 200
    accounts = accounts_res.json()
    assert len(accounts) == 5, f"Expected 5 platforms, got {len(accounts)}"
    account_platforms = [a["platform"] for a in accounts]
    assert account_platforms == ["youtube", "instagram", "facebook", "x", "linkedin"], f"Unexpected platforms: {account_platforms}"
    assert "tiktok" not in account_platforms
    print(f"[PASS] 8a. Social Accounts Endpoint returns exact 5 platforms: {account_platforms}")

    # Platforms live status endpoint
    status_res = client.get("/api/v1/social/platforms/status", headers=auth_headers)
    assert status_res.status_code == 200
    p_status = status_res.json()
    assert p_status["total_supported"] == 5
    print("[PASS] 8b. Social Platforms Live Connector Status Endpoint verified")

    # Sync YouTube
    yt_acc = next(a for a in accounts if a["platform"] == "youtube")
    yt_sync = client.post(f"/api/v1/social/accounts/{yt_acc['id']}/sync", headers=auth_headers)
    assert yt_sync.status_code == 200
    assert yt_sync.json()["is_connected"] is True
    print("[PASS] 8c. YouTube Sync verified")

    # Sync Instagram (clean non-crashing configuration response)
    ig_acc = next(a for a in accounts if a["platform"] == "instagram")
    ig_sync = client.post(f"/api/v1/social/accounts/{ig_acc['id']}/sync", headers=auth_headers)
    assert ig_sync.status_code == 200
    print("[PASS] 8d. Instagram Sync clean behavior verified")

    # 9. Individual Service Unit Inspections
    yt_svc = YouTubeIntegrationService()
    yt_info = yt_svc.fetch_public_channel_data()
    assert yt_info["subscribers"] == 1420000
    print(f"[PASS] 9a. YouTubeIntegrationService: {yt_info['status']}")

    ig_svc = InstagramIntegrationService()
    ig_info = ig_svc.fetch_profile_and_posts()
    assert ig_info["status"] == "Configuration Required / Not Connected"
    print(f"[PASS] 9b. InstagramIntegrationService (RapidAPI): {ig_info['status']}")

    fb_svc = FacebookIntegrationService()
    fb_status = fb_svc.get_status()
    assert fb_status["status"] == "Configuration Required / Not Connected"
    print(f"[PASS] 9c. FacebookIntegrationService: {fb_status['status']}")

    x_svc = XIntegrationService()
    x_status = x_svc.get_status()
    assert x_status["status"] == "Configuration Required / Not Connected"
    print(f"[PASS] 9d. XIntegrationService: {x_status['status']}")

    li_svc = LinkedInIntegrationService()
    li_status = li_svc.get_status()
    assert li_status["status"] == "Configuration Required / Not Connected"
    print(f"[PASS] 9e. LinkedInIntegrationService: {li_status['status']}")

    # Restore clean state
    from app.core.database import SessionLocal
    from app.db.seed import seed_database
    db_clean = SessionLocal()
    try:
        seed_database(db_clean, force=True)
    finally:
        db_clean.close()

    print("==================================================")
    print("ALL MILESTONE 3 VERIFICATION TESTS PASSED (100%)!")
    print("==================================================")

if __name__ == "__main__":
    run_tests()

