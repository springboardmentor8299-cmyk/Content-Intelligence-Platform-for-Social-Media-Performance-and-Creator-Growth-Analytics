
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from sqlalchemy.orm import Session

from app.auth.security import get_current_creator, get_current_user
from app.database import get_db
from app.models import User
from app.models.earnings import Earnings
from app.models.sponsorship import Sponsorship
from app.models.analytics import Analytics

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

router = APIRouter(
    prefix="/reports",
    tags=["Reports"],
    dependencies=[Depends(get_current_creator)],
)


def create_excel_response(workbook, filename: str):
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"'
        }
    )


def style_header(sheet):
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78"
        )
        cell.alignment = Alignment(horizontal="center")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for column in sheet.columns:
        max_length = max(
            len(str(cell.value or ""))
            for cell in column
        )
        sheet.column_dimensions[
            column[0].column_letter
        ].width = min(max_length + 3, 35)


# =====================================================
# EXPORT EARNINGS TO EXCEL
# =====================================================

@router.get("/earnings/excel")
def export_earnings_excel(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    earnings = (
        db.query(Earnings)
        .filter(Earnings.creator_id == current_user.id)
        .order_by(Earnings.date.asc())
        .all()
    )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Earnings"

    sheet.append([
        "Date",
        "Estimated Revenue",
        "Ad Revenue",
        "Monetized Views",
        "Currency"
    ])

    for item in earnings:
        sheet.append([
            item.date.isoformat(),
            float(item.estimated_revenue or 0),
            float(item.ad_revenue or 0),
            int(item.monetized_views or 0),
            item.currency or "USD"
        ])

    if not earnings:
        sheet.append(["No earnings records found"])

    style_header(sheet)

    return create_excel_response(
        workbook,
        "creatoriq_earnings.xlsx"
    )


# =====================================================
# EXPORT SPONSORSHIPS TO EXCEL
# =====================================================

@router.get("/sponsorships/excel")
def export_sponsorships_excel(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    sponsorships = (
        db.query(Sponsorship)
        .filter(Sponsorship.creator_id == current_user.id)
        .order_by(Sponsorship.id.asc())
        .all()
    )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sponsorships"

    sheet.append([
        "Brand",
        "Campaign",
        "Amount",
        "Currency",
        "Payment Status",
        "Start Date",
        "Due Date",
        "Paid Date",
        "Notes"
    ])

    for item in sponsorships:
        sheet.append([
            item.brand_name,
            item.campaign_name,
            float(item.amount or 0),
            item.currency or "INR",
            item.payment_status,
            item.start_date.isoformat() if item.start_date else "",
            item.due_date.isoformat() if item.due_date else "",
            item.paid_date.isoformat() if item.paid_date else "",
            item.notes or ""
        ])

    if not sponsorships:
        sheet.append(["No sponsorship records found"])

    style_header(sheet)

    return create_excel_response(
        workbook,
        "creatoriq_sponsorships.xlsx"
    )


def create_pdf_response(title, headers, rows, filename):
    buffer = BytesIO()

    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=25,
        leftMargin=25,
        topMargin=30,
        bottomMargin=30,
    )

    styles = getSampleStyleSheet()
    story = [
        Paragraph(title, styles["Title"]),
        Spacer(1, 16),
    ]

    # Make long text wrap inside table cells.
    cell_style = ParagraphStyle(
        "TableCell",
        parent=styles["BodyText"],
        fontSize=7,
        leading=9,
    )

    table_data = [
        [Paragraph(str(header), cell_style) for header in headers]
    ]

    for row in rows:
        table_data.append([
            Paragraph(str(value if value is not None else ""), cell_style)
            for value in row
        ])

    if not rows:
        table_data.append([
            Paragraph("No records found.", cell_style)
        ])

    table = Table(table_data, repeatRows=1)

    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1),
         [colors.white, colors.HexColor("#EAF2F8")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))

    story.append(table)
    document.build(story)

    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"'
        },
    )


@router.get("/earnings/pdf")
def export_earnings_pdf(
    db=Depends(get_db),
    current_user=Depends(get_current_user),
):
    earnings = (
        db.query(Earnings)
        .filter(Earnings.creator_id == current_user.id)
        .order_by(Earnings.date.asc())
        .all()
    )

    rows = [
        [
            item.date.isoformat(),
            float(item.estimated_revenue or 0),
            float(item.ad_revenue or 0),
            int(item.monetized_views or 0),
            item.currency or "USD",
        ]
        for item in earnings
    ]

    return create_pdf_response(
        "CreatorIQ - Earnings Report",
        ["Date", "Estimated Revenue", "Ad Revenue",
         "Monetized Views", "Currency"],
        rows,
        "creatoriq_earnings.pdf",
    )


@router.get("/sponsorships/pdf")
def export_sponsorships_pdf(
    db=Depends(get_db),
    current_user=Depends(get_current_user),
):
    sponsorships = (
        db.query(Sponsorship)
        .filter(Sponsorship.creator_id == current_user.id)
        .order_by(Sponsorship.id.asc())
        .all()
    )

    rows = [
        [
            item.brand_name,
            item.campaign_name,
            float(item.amount or 0),
            item.currency or "INR",
            item.payment_status,
            item.start_date.isoformat() if item.start_date else "",
            item.due_date.isoformat() if item.due_date else "",
            item.paid_date.isoformat() if item.paid_date else "",
        ]
        for item in sponsorships
    ]

    return create_pdf_response(
        "CreatorIQ - Sponsorship Report",
        ["Brand", "Campaign", "Amount", "Currency",
         "Status", "Start Date", "Due Date", "Paid Date"],
        rows,
        "creatoriq_sponsorships.pdf",
    )


@router.get("/analytics/summary")
def get_combined_analytics_summary(
    db=Depends(get_db),
    current_user=Depends(get_current_user),
):
    creator_id = current_user.id

    analytics_records = (
        db.query(Analytics)
        .filter(Analytics.creator_id == creator_id)
        .all()
    )

    earnings_records = (
        db.query(Earnings)
        .filter(Earnings.creator_id == creator_id)
        .all()
    )

    sponsorship_records = (
        db.query(Sponsorship)
        .filter(Sponsorship.creator_id == creator_id)
        .all()
    )

    total_views = sum(item.views or 0 for item in analytics_records)
    total_likes = sum(item.likes or 0 for item in analytics_records)
    total_comments = sum(item.comments or 0 for item in analytics_records)
    total_shares = sum(item.shares or 0 for item in analytics_records)
    total_saves = sum(item.saves or 0 for item in analytics_records)

    total_watch_time = sum(
        item.watch_time or 0 for item in analytics_records
    )

    latest_analytics = max(
        analytics_records,
        key=lambda item: item.date,
        default=None,
    )

    total_estimated_revenue = sum(
        item.estimated_revenue or 0 for item in earnings_records
    )

    total_ad_revenue = sum(
        item.ad_revenue or 0 for item in earnings_records
    )

    total_sponsorship_value = sum(
        item.amount or 0 for item in sponsorship_records
    )

    paid_sponsorship_value = sum(
        item.amount or 0
        for item in sponsorship_records
        if (item.payment_status or "").lower() == "paid"
    )

    pending_sponsorship_value = sum(
        item.amount or 0
        for item in sponsorship_records
        if (item.payment_status or "").lower() == "pending"
    )

    return {
        "analytics": {
            "records": len(analytics_records),
            "total_views": total_views,
            "total_likes": total_likes,
            "total_comments": total_comments,
            "total_shares": total_shares,
            "total_saves": total_saves,
            "total_watch_time": total_watch_time,
            "latest_followers": (
                latest_analytics.followers if latest_analytics else 0
            ),
            "latest_reach": (
                latest_analytics.reach if latest_analytics else 0
            ),
            "latest_engagement_rate": (
                latest_analytics.engagement_rate
                if latest_analytics else 0
            ),
        },
        "earnings": {
            "records": len(earnings_records),
            "total_estimated_revenue": total_estimated_revenue,
            "total_ad_revenue": total_ad_revenue,
            "currency_note": (
                "Revenue totals are raw sums; ensure records use "
                "the same currency before interpreting totals."
            ),
        },
        "sponsorships": {
            "records": len(sponsorship_records),
            "total_value": total_sponsorship_value,
            "paid_value": paid_sponsorship_value,
            "pending_value": pending_sponsorship_value,
            "currency_note": (
                "Sponsorship totals are raw sums; ensure records use "
                "the same currency before interpreting totals."
            ),
        },
    }


@router.get("/combined/excel")
def export_combined_analytics_excel(
    db=Depends(get_db),
    current_user=Depends(get_current_user),
):
    creator_id = current_user.id

    # Fetch only the logged-in creator's data
    analytics_records = (
        db.query(Analytics)
        .filter(Analytics.creator_id == creator_id)
        .order_by(Analytics.date.asc())
        .all()
    )

    earnings_records = (
        db.query(Earnings)
        .filter(Earnings.creator_id == creator_id)
        .order_by(Earnings.date.asc())
        .all()
    )

    sponsorship_records = (
        db.query(Sponsorship)
        .filter(Sponsorship.creator_id == creator_id)
        .order_by(Sponsorship.id.asc())
        .all()
    )

    workbook = Workbook()

    # Sheet 1: Overview
    overview = workbook.active
    overview.title = "Overview"

    total_estimated_revenue = sum(
        item.estimated_revenue or 0 for item in earnings_records
    )
    total_ad_revenue = sum(
        item.ad_revenue or 0 for item in earnings_records
    )
    total_sponsorship_value = sum(
        item.amount or 0 for item in sponsorship_records
    )
    paid_sponsorship_value = sum(
        item.amount or 0
        for item in sponsorship_records
        if (item.payment_status or "").lower() == "paid"
    )
    pending_sponsorship_value = sum(
        item.amount or 0
        for item in sponsorship_records
        if (item.payment_status or "").lower() == "pending"
    )

    latest_analytics = max(
        analytics_records,
        key=lambda item: item.date,
        default=None,
    )

    overview.append(["Metric", "Value"])
    overview.append(["Analytics Records", len(analytics_records)])
    overview.append(["Earnings Records", len(earnings_records)])
    overview.append(["Sponsorship Records", len(sponsorship_records)])
    overview.append(["Total Views", sum(x.views or 0 for x in analytics_records)])
    overview.append(["Total Likes", sum(x.likes or 0 for x in analytics_records)])
    overview.append(["Total Comments", sum(x.comments or 0 for x in analytics_records)])
    overview.append(["Estimated Revenue Total", total_estimated_revenue])
    overview.append(["Ad Revenue Total", total_ad_revenue])
    overview.append(["Sponsorship Value Total", total_sponsorship_value])
    overview.append(["Paid Sponsorship Value", paid_sponsorship_value])
    overview.append(["Pending Sponsorship Value", pending_sponsorship_value])
    overview.append([
        "Latest Followers",
        latest_analytics.followers if latest_analytics else 0,
    ])
    overview.append([
        "Latest Reach",
        latest_analytics.reach if latest_analytics else 0,
    ])
    overview.append([
        "Latest Engagement Rate",
        latest_analytics.engagement_rate if latest_analytics else 0,
    ])
    overview.append([
        "Currency Note",
        "Revenue and sponsorship totals are separate raw sums. "
        "Use consistent currencies before comparing or adding them.",
    ])

    # Sheet 2: Analytics details
    sheet = workbook.create_sheet("Analytics")
    sheet.append([
        "Date", "Views", "Likes", "Comments", "Shares", "Saves",
        "Watch Time", "Followers", "Reach", "Engagement Rate",
    ])

    for item in analytics_records:
        sheet.append([
            item.date.isoformat(),
            item.views or 0,
            item.likes or 0,
            item.comments or 0,
            item.shares or 0,
            item.saves or 0,
            item.watch_time or 0,
            item.followers or 0,
            item.reach or 0,
            item.engagement_rate or 0,
        ])

    if not analytics_records:
        sheet.append(["No analytics records found"])

    # Sheet 3: Earnings details
    sheet = workbook.create_sheet("Earnings")
    sheet.append([
        "Date", "Estimated Revenue", "Ad Revenue",
        "Monetized Views", "Currency",
    ])

    for item in earnings_records:
        sheet.append([
            item.date.isoformat(),
            float(item.estimated_revenue or 0),
            float(item.ad_revenue or 0),
            item.monetized_views or 0,
            item.currency or "USD",
        ])

    if not earnings_records:
        sheet.append(["No earnings records found"])

    # Sheet 4: Sponsorship details
    sheet = workbook.create_sheet("Sponsorships")
    sheet.append([
        "Brand", "Campaign", "Amount", "Currency", "Payment Status",
        "Start Date", "Due Date", "Paid Date", "Notes",
    ])

    for item in sponsorship_records:
        sheet.append([
            item.brand_name,
            item.campaign_name,
            float(item.amount or 0),
            item.currency or "INR",
            item.payment_status,
            item.start_date.isoformat() if item.start_date else "",
            item.due_date.isoformat() if item.due_date else "",
            item.paid_date.isoformat() if item.paid_date else "",
            item.notes or "",
        ])

    if not sponsorship_records:
        sheet.append(["No sponsorship records found"])

    # Apply the existing formatting helper to every sheet
    for worksheet in workbook.worksheets:
        style_header(worksheet)

    return create_excel_response(
        workbook,
        "creatoriq_combined_report.xlsx",
    )


@router.get("/combined/pdf")
def export_combined_analytics_pdf(
    db=Depends(get_db),
    current_user=Depends(get_current_user),
):
    creator_id = current_user.id

    analytics_records = (
        db.query(Analytics)
        .filter(Analytics.creator_id == creator_id)
        .order_by(Analytics.date.asc())
        .all()
    )

    earnings_records = (
        db.query(Earnings)
        .filter(Earnings.creator_id == creator_id)
        .order_by(Earnings.date.asc())
        .all()
    )

    sponsorship_records = (
        db.query(Sponsorship)
        .filter(Sponsorship.creator_id == creator_id)
        .order_by(Sponsorship.id.asc())
        .all()
    )

    total_views = sum(x.views or 0 for x in analytics_records)
    total_likes = sum(x.likes or 0 for x in analytics_records)
    total_comments = sum(x.comments or 0 for x in analytics_records)

    total_estimated_revenue = sum(
        x.estimated_revenue or 0 for x in earnings_records
    )
    total_ad_revenue = sum(
        x.ad_revenue or 0 for x in earnings_records
    )

    total_sponsorship_value = sum(
        x.amount or 0 for x in sponsorship_records
    )
    paid_sponsorship_value = sum(
        x.amount or 0 for x in sponsorship_records
        if (x.payment_status or "").lower() == "paid"
    )
    pending_sponsorship_value = sum(
        x.amount or 0 for x in sponsorship_records
        if (x.payment_status or "").lower() == "pending"
    )

    latest_analytics = max(
        analytics_records,
        key=lambda x: x.date,
        default=None,
    )

    rows = [
        ["Category", "Metric", "Value"],
        ["Analytics", "Number of records", len(analytics_records)],
        ["Analytics", "Total views", total_views],
        ["Analytics", "Total likes", total_likes],
        ["Analytics", "Total comments", total_comments],
        [
            "Analytics",
            "Latest followers",
            latest_analytics.followers if latest_analytics else 0,
        ],
        [
            "Analytics",
            "Latest reach",
            latest_analytics.reach if latest_analytics else 0,
        ],
        [
            "Analytics",
            "Latest engagement rate",
            latest_analytics.engagement_rate if latest_analytics else 0,
        ],
        ["Earnings", "Number of records", len(earnings_records)],
        ["Earnings", "Estimated revenue total", total_estimated_revenue],
        ["Earnings", "Ad revenue total", total_ad_revenue],
        ["Sponsorships", "Number of records", len(sponsorship_records)],
        ["Sponsorships", "Total sponsorship value", total_sponsorship_value],
        ["Sponsorships", "Paid sponsorship value", paid_sponsorship_value],
        [
            "Sponsorships",
            "Pending sponsorship value",
            pending_sponsorship_value,
        ],
    ]

    return create_pdf_response(
        "CreatorIQ - Combined Analytics Report",
        rows[0],
        rows[1:],
        "creatoriq_combined_report.pdf",
    )


@router.get("/audience")
def get_audience_report(
    db=Depends(get_db),
    current_user=Depends(get_current_user),
):
    records = (
        db.query(Analytics)
        .filter(Analytics.creator_id == current_user.id)
        .order_by(Analytics.date.asc(), Analytics.id.asc())
        .all()
    )

    return {
        "total_records": len(records),
        "audience_trend": [
            {
                "date": item.date.isoformat(),
                "followers": item.followers or 0,
                "reach": item.reach or 0,
            }
            for item in records
        ],
    }


@router.get("/growth")
def get_growth_report(
    db=Depends(get_db),
    current_user=Depends(get_current_user),
):
    records = (
        db.query(Analytics)
        .filter(Analytics.creator_id == current_user.id)
        .order_by(Analytics.date.asc(), Analytics.id.asc())
        .all()
    )

    growth_trend = []

    for index, item in enumerate(records):
        previous = records[index - 1] if index > 0 else None

        previous_followers = (
            previous.followers or 0 if previous else 0
        )
        previous_views = previous.views or 0 if previous else 0

        current_followers = item.followers or 0
        current_views = item.views or 0

        follower_change = (
            current_followers - previous_followers
            if previous else None
        )
        views_change = (
            current_views - previous_views
            if previous else None
        )

        follower_growth_percent = (
            round(follower_change / previous_followers * 100, 2)
            if previous and previous_followers > 0
            else None
        )

        views_growth_percent = (
            round(views_change / previous_views * 100, 2)
            if previous and previous_views > 0
            else None
        )

        growth_trend.append({
            "date": item.date.isoformat(),
            "followers": current_followers,
            "views": current_views,
            "follower_change": follower_change,
            "views_change": views_change,
            "follower_growth_percent": follower_growth_percent,
            "views_growth_percent": views_growth_percent,
            "engagement_rate": item.engagement_rate or 0,
        })

    return {
        "total_records": len(records),
        "growth_trend": growth_trend,
        "note": (
            "Growth compares each record with the previous record. "
            "The first record has no previous comparison. "
            "Meaningful growth analysis requires records from "
            "different dates."
        ),
    }
